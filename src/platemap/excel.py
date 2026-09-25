"""Excel workbook writer/reader for plate layouts and reader data."""

import math
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from platemap.layout import LAYOUT_COLUMNS, ROLE_FILL, LayoutRow
from platemap.wells import COLS_FULL, ROWS_FULL

MAPPED_COLUMNS = (
    "well",
    "role",
    "short_id",
    "label",
    "conc_ugml",
    "sample_name",
    "dilution_factor",
    "replicate",
    "absorbance",
)

_READER_SHEET_RE = re.compile(r"^Plate (\d+) Reader$")


def _write_grid_headers(ws: Worksheet) -> None:
    """Write the 1..12 and A..H headers used by the Map and Reader sheets."""
    for i, col in enumerate(COLS_FULL, start=2):
        ws.cell(row=1, column=i, value=col)
    for i, row in enumerate(ROWS_FULL, start=2):
        ws.cell(row=i, column=1, value=row)


def write_excel(rows: list[LayoutRow], path: str, experiment: str = "", date: str = "") -> None:
    """Write the Info, Layout, and per-plate Map/Reader/Mapped sheets."""
    wb = Workbook()
    wb.remove(wb.active)

    info = wb.create_sheet("Info")
    info["A1"] = "Experiment"
    info["B1"] = experiment
    info["A2"] = "Date"
    info["B2"] = date

    layout = wb.create_sheet("Layout")
    for i, col in enumerate(LAYOUT_COLUMNS, start=1):
        cell = layout.cell(row=1, column=i, value=col)
        cell.font = Font(bold=True)
    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(LAYOUT_COLUMNS, start=1):
            value = getattr(row, col)
            layout.cell(row=r, column=c, value="" if value is None else value)

    plates = sorted({row.plate for row in rows})
    row_index = {r: i for i, r in enumerate(ROWS_FULL, start=2)}

    for plate in plates:
        plate_rows = [row for row in rows if row.plate == plate]

        map_ws = wb.create_sheet(f"Plate {plate} Map")
        _write_grid_headers(map_ws)
        for row in plate_rows:
            cell = map_ws.cell(row=row_index[row.row], column=row.col + 1, value=row.label)
            fill_hex = ROLE_FILL[row.role]
            cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")

        reader_ws = wb.create_sheet(f"Plate {plate} Reader")
        _write_grid_headers(reader_ws)

        mapped_ws = wb.create_sheet(f"Plate {plate} Mapped")
        for i, col in enumerate(MAPPED_COLUMNS, start=1):
            cell = mapped_ws.cell(row=1, column=i, value=col)
            cell.font = Font(bold=True)
        for r, row in enumerate(plate_rows, start=2):
            values = {
                "well": row.well,
                "role": row.role,
                "short_id": row.short_id,
                "label": row.label,
                "conc_ugml": "" if row.conc_ugml is None else row.conc_ugml,
                "sample_name": "" if row.sample_name is None else row.sample_name,
                "dilution_factor": "" if row.dilution_factor is None else row.dilution_factor,
                "replicate": row.replicate,
                "absorbance": (
                    f"=INDEX('Plate {plate} Reader'!$B$2:$M$9,{row_index[row.row] - 1},{row.col})"
                ),
            }
            for c, col in enumerate(MAPPED_COLUMNS, start=1):
                mapped_ws.cell(row=r, column=c, value=values[col])

    wb.save(path)


def read_layout(workbook) -> list[LayoutRow]:
    """Read the Layout sheet of a workbook (path or Workbook) into LayoutRow objects."""
    wb = load_workbook(workbook) if not isinstance(workbook, Workbook) else workbook
    ws = wb["Layout"]

    rows: list[LayoutRow] = []
    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None or str(v).strip() == "" for v in row_cells):
            continue
        values = dict(zip(LAYOUT_COLUMNS, row_cells))
        rows.append(
            LayoutRow(
                plate=int(values["plate"]),
                well=str(values["well"]),
                row=str(values["row"]),
                col=int(values["col"]),
                role=str(values["role"]),
                short_id=str(values["short_id"]),
                label=str(values["label"]),
                conc_ugml=None if values["conc_ugml"] in (None, "") else float(values["conc_ugml"]),
                sample_name=None if values["sample_name"] in (None, "") else str(values["sample_name"]),
                dilution_factor=(
                    None
                    if values["dilution_factor"] in (None, "")
                    else float(values["dilution_factor"])
                ),
                replicate=int(values["replicate"]),
                notes="" if values["notes"] is None else str(values["notes"]),
            )
        )
    return rows


def plate_count(workbook) -> int:
    """Count the "Plate N Reader" sheets in a workbook (path or Workbook)."""
    wb = load_workbook(workbook) if not isinstance(workbook, Workbook) else workbook
    return sum(1 for name in wb.sheetnames if _READER_SHEET_RE.match(name))


def fill_reader(workbook, plate_values: dict[int, dict[str, float]], out_path: str) -> Path:
    """Write reader values into the Plate N Reader grids; NaN leaves the cell empty."""
    wb = load_workbook(workbook) if not isinstance(workbook, Workbook) else workbook
    row_index = {r: i for i, r in enumerate(ROWS_FULL, start=2)}

    for plate, values in plate_values.items():
        ws = wb[f"Plate {plate} Reader"]
        for well, value in values.items():
            if value is None or (isinstance(value, float) and math.isnan(value)):
                continue
            row_letter, col = well[0], int(well[1:])
            ws.cell(row=row_index[row_letter], column=col + 1, value=value)

    wb.save(out_path)
    return Path(out_path)


def compute_mapped_rows(
    rows: list[LayoutRow], plate_values: dict[int, dict[str, float]]
) -> list[dict]:
    """Combine layout rows with reader values into mapped-row dicts."""
    mapped = []
    for row in rows:
        value = plate_values.get(row.plate, {}).get(row.well, math.nan)
        mapped.append(
            {
                "plate": row.plate,
                "well": row.well,
                "role": row.role,
                "short_id": row.short_id,
                "label": row.label,
                "conc_ugml": row.conc_ugml,
                "sample_name": row.sample_name,
                "dilution_factor": row.dilution_factor,
                "replicate": row.replicate,
                "absorbance": value,
            }
        )
    return mapped
