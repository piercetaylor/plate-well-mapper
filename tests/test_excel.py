import math

from openpyxl import load_workbook

from platemap.excel import compute_mapped_rows, fill_reader, plate_count, read_layout, write_excel
from platemap.layout import build_layout
from platemap.samples import Sample


def _samples(n):
    return [Sample(sample_name=f"Samp{i}") for i in range(1, n + 1)]


def test_write_excel_sheet_names(tmp_path):
    rows = build_layout(_samples(1), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path), experiment="Exp1", date="2026-01-01")

    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Info",
        "Layout",
        "Plate 1 Map",
        "Plate 1 Reader",
        "Plate 1 Mapped",
    ]
    assert wb["Info"]["A1"].value == "Experiment"
    assert wb["Info"]["B1"].value == "Exp1"
    assert wb["Info"]["A2"].value == "Date"
    assert wb["Info"]["B2"].value == "2026-01-01"
    assert wb["Layout"]["A1"].value == "plate"
    assert wb["Layout"]["A1"].font.bold is True


def test_reader_headers_and_empty_grid(tmp_path):
    rows = build_layout(_samples(1), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path))
    wb = load_workbook(path)
    reader = wb["Plate 1 Reader"]
    assert [reader.cell(row=1, column=c).value for c in range(2, 14)] == list(range(1, 13))
    assert [reader.cell(row=r, column=1).value for r in range(2, 10)] == list("ABCDEFGH")
    assert reader["B2"].value is None
    assert reader["M9"].value is None


def test_mapped_formula(tmp_path):
    rows = build_layout(_samples(27), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path))
    wb = load_workbook(path)
    mapped = wb["Plate 1 Mapped"]
    b7_row = next(
        r for r in range(2, mapped.max_row + 1) if mapped.cell(row=r, column=1).value == "B7"
    )
    formula = mapped.cell(row=b7_row, column=9).value
    assert formula == "=INDEX('Plate 1 Reader'!$B$2:$M$9,2,7)"


def test_read_layout_roundtrip(tmp_path):
    rows = build_layout(_samples(27), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path))
    roundtrip = read_layout(str(path))
    assert len(roundtrip) == len(rows)
    assert roundtrip[0].well == rows[0].well
    assert roundtrip[0].conc_ugml == rows[0].conc_ugml


def test_plate_count(tmp_path):
    rows = build_layout(_samples(27), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path))
    assert plate_count(str(path)) == 2


def test_fill_reader_and_compute_mapped_rows(tmp_path):
    rows = build_layout(_samples(1), avoid_edges=False)
    path = tmp_path / "out.xlsx"
    write_excel(rows, str(path))

    out_path = tmp_path / "filled.xlsx"
    plate_values = {1: {"A1": 0.5, "A2": float("nan")}}
    result = fill_reader(str(path), plate_values, str(out_path))
    assert result == out_path

    wb = load_workbook(out_path)
    reader = wb["Plate 1 Reader"]
    assert reader["B2"].value == 0.5
    assert reader["C2"].value is None

    mapped = compute_mapped_rows(rows, plate_values)
    a1 = next(m for m in mapped if m["well"] == "A1")
    assert a1["absorbance"] == 0.5
    b1 = next(m for m in mapped if m["well"] == "B1")
    assert math.isnan(b1["absorbance"])
